"""קישור פרויקטים — the completion table (Excel, option ב, 2026-10-19). Pure, no DB."""
import io

import pytest
from openpyxl import load_workbook

from app.services import lessons_link_sheet as sheet
from app.services import lessons_project_link as lpl

STAGES = lpl.STAGE_OPTIONS
TYPES = lpl.TYPE_LABELS
USERS = [{"id": 7, "username": "משה כהן"}, {"id": 9, "username": "דנה לוי "}]

ITEMS = [
    {"key": "stage|P-3", "kind": "stage", "identifier": "P-3", "name": "תחנה ג", "manager": "יוסי",
     "current": "שלב מוזר", "note": "השלב בקובץ לא מוכר"},
    {"key": "project_type|P-4", "kind": "project_type", "identifier": "P-4", "name": "תחנה ד", "manager": "",
     "current": "משהו", "note": ""},
    {"key": "alias|יוסי", "kind": "alias", "alias": "יוסי ", "identifier": "", "name": "3 פרויקטים",
     "manager": "יוסי", "current": "", "note": ""},
    {"key": "pm|12", "kind": "pm", "identifier": "P-8", "name": "תחנה ח", "manager": "", "current": "",
     "project_id": 12, "note": ""},
]


def _fill(data: bytes, values: dict[str, str], current: dict[str, str] | None = None) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    ws = wb[sheet.SHEET_ITEMS]
    for row in ws.iter_rows(min_row=2):
        key = row[sheet.COL["מפתח"]].value
        if key in values:
            row[sheet.COL["ערך להשלמה"]].value = values[key]
        if current and key in current:
            row[sheet.COL["ערך קיים"]].value = current[key]
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _open():
    return {i["key"]: i for i in ITEMS}


def test_export_has_one_row_per_item_with_dropdowns_and_a_hidden_key():
    data = sheet.build_workbook(ITEMS, STAGES, TYPES, USERS)
    wb = load_workbook(io.BytesIO(data))
    ws = wb[sheet.SHEET_ITEMS]
    assert tuple(c.value for c in ws[1]) == sheet.HEADERS
    assert ws.max_row == len(ITEMS) + 1
    assert ws.column_dimensions["A"].hidden
    assert wb[sheet.SHEET_LISTS].sheet_state == "hidden"
    assert len(ws.data_validations.dataValidation) == 3          # stage / type / user lists
    # the user list carries the id — that is what is read back, never the name
    users = [c.value for c in wb[sheet.SHEET_LISTS]["C"][1:] if c.value]
    assert users == ["משה כהן · #7", "דנה לוי · #9"]


def test_round_trip_reads_back_every_filled_row():
    data = _fill(sheet.build_workbook(ITEMS, STAGES, TYPES, USERS),
                 {"stage|P-3": "בחירת קבלן", "project_type|P-4": "ניידות", "alias|יוסי": "משה כהן · #7",
                  "pm|12": "דנה לוי · #9"})
    check = sheet.validate(sheet.parse_workbook(data), _open(), STAGES, TYPES, {7, 9})
    assert check["counts"]["ok"] == 4
    got = {a["key"]: a["value"] for a in check["apply"]}
    assert got == {"stage|P-3": "בחירת קבלן", "project_type|P-4": "mobility", "alias|יוסי": 7, "pm|12": 9}
    # the raw file name travels with the alias row (the file spells it with a space)
    assert next(a for a in check["apply"] if a["kind"] == "alias")["alias"] == "יוסי "


def test_an_empty_cell_changes_nothing():
    data = sheet.build_workbook(ITEMS, STAGES, TYPES, USERS)
    check = sheet.validate(sheet.parse_workbook(data), _open(), STAGES, TYPES, {7, 9})
    assert check["apply"] == [] and check["rows"] == [] and check["counts"]["empty"] == len(ITEMS)


@pytest.mark.parametrize(("value", "status"), [
    ("שלב שאין", "invalid"),
    ("בחירת קבלן", "ok"),
])
def test_stage_must_be_a_legal_value(value, status):
    data = _fill(sheet.build_workbook(ITEMS, STAGES, TYPES, USERS), {"stage|P-3": value})
    assert sheet.validate(sheet.parse_workbook(data), _open(), STAGES, TYPES, {7})["rows"][0]["status"] == status


def test_a_user_is_read_by_id_only_and_must_exist():
    data = _fill(sheet.build_workbook(ITEMS, STAGES, TYPES, USERS),
                 {"alias|יוסי": "משה כהן", "pm|12": "מישהו · #55"})
    rows = sheet.validate(sheet.parse_workbook(data), _open(), STAGES, TYPES, {7, 9})["rows"]
    assert [r["status"] for r in rows] == ["invalid", "invalid"]


def test_a_row_whose_current_value_changed_since_export_is_stale():
    data = _fill(sheet.build_workbook(ITEMS, STAGES, TYPES, USERS), {"stage|P-3": "בחירת קבלן"})
    now = {**_open(), "stage|P-3": {**_open()["stage|P-3"], "current": "שלב אחר"}}
    assert sheet.validate(sheet.parse_workbook(data), now, STAGES, TYPES, {7})["rows"][0]["status"] == "stale"


def test_an_item_handled_since_export_is_skipped():
    data = _fill(sheet.build_workbook(ITEMS, STAGES, TYPES, USERS), {"stage|P-3": "בחירת קבלן"})
    now = {k: v for k, v in _open().items() if k != "stage|P-3"}
    check = sheet.validate(sheet.parse_workbook(data), now, STAGES, TYPES, {7})
    assert check["rows"][0]["status"] == "closed" and check["apply"] == []


def test_not_the_completion_file_is_refused():
    with pytest.raises(sheet.SheetError):
        sheet.parse_workbook(b"not a workbook")
    from openpyxl import Workbook
    out = io.BytesIO()
    Workbook().save(out)
    with pytest.raises(sheet.SheetError):
        sheet.parse_workbook(out.getvalue())


# ------------------------------------------------------------------ what the table asks for

def _payload(held=(), projects=()):
    return {"summary": {"held": list(held)}, "projects": list(projects)}


def test_completion_items_ask_only_what_a_person_can_give():
    held = [{"identifier": "P-3", "name": "ג", "manager": "יוסי", "stage": "שלב מוזר", "project_type": "הקמה",
             "missing": ["stage"]}]
    projects = [
        {"id": 12, "name": "ח", "identifier": "P-8", "issues": ["manager_placeholder"],
         "manager_check": {"status": "no_name"}},
        # a manager that DIFFERS is an approval, not a completion — not in the table
        {"id": 13, "name": "ט", "identifier": "P-9", "issues": ["manager_differs"],
         "manager_check": {"status": "differs"}},
        # handled — not in the table
        {"id": 14, "name": "י", "identifier": "P-10", "issues": [], "manager_check": {"status": "match"}},
    ]
    managers = [{"name": "יוסי ", "projects": 3, "problem": "alias_unlinked", "username": None},
                {"name": "משה כהן", "projects": 5, "problem": None, "username": "משה כהן"}]
    items = lpl.completion_items(_payload(held, projects), managers)
    assert [i["key"] for i in items] == ["stage|P-3", "alias|יוסי", "pm|12"]
    assert items[0]["current"] == "שלב מוזר"
    assert items[1]["alias"] == "יוסי "


def test_completion_routes_are_registered_before_the_gateway():
    from app.main import app
    paths = [getattr(r, "path", "") for r in app.routes]
    gw = next(i for i, p in enumerate(paths) if p.startswith("/lessons/api/{"))
    assert paths.index("/lessons/api/_project_links/completion.xlsx") < gw
    assert paths.index("/lessons/api/_project_links/completion") < gw
