"""The master file's "סטטוס התקשרויות קבלנים ותיחורים" sheet → its projects page.
Synthetic rows only — the repo is public and the real file is inside information.
"""
import datetime
import re
from pathlib import Path

from app.services import contracts_sheet as cs
from app.services import il_format

D = datetime.datetime
REPORT = datetime.date(2026, 10, 14)
HEADER = ["מספר הפרויקט", None, "סוג הליך", "מנהל פרויקט", "תאריך העברה לרכש", "סטטוס",
          "צפי קבלת היתר \\ יעד עליה לקרקע", "תאריך חישמול ת.פ.", "תאריך חשמול מסתמן", "9.7.26"]


PLAN, FORECAST = D(2027, 3, 30), D(2027, 10, 30)


def row(ident, name, status, *, proc="מכרז פומבי", transfer="05.02.2026", ground="צפי עליה לקרקע 30/09/2026",
        plan=PLAN, fc=FORECAST):
    return [ident, name, proc, "מנהל א", transfer, status, ground, plan, fc, None]


ROWS = [
    HEADER,
    row("X-101", "תחנה א X-101", "התקיים סיור ב10.06 הוגשו הצעות ב06.07 . צפי למכתב זכייה עד 15.07"),
    row("X-102", "תחנה ב X-102", "ב11.05 אישור קבלן בועדה .ב17.05 התקבל מכתב זכייה ע\"ש קבלן אלפא בסך 1,200,000 ש\"ח "
        "מס הזמנה 4001000001 ב11.06 התקבלה הזמנה חתומה. הסתיים", ground="התקבל היתר ב15/03/2026 , צפי עליה לקרקע 01/06/2026"),
    row("X-103", "תחנה ג X-103", "נקבע סיור ל10.06 . ב07.06 מנ\"פ ביקש לבטל את סיור הקבלנים", plan=D(2026, 11, 30)),
    row("X-104", "תחנה ד X-104", ".. הגשת הצעות נדחתה ל06.07 ואח\"כ שוב נדחתה ל13.7"),
    row(None, "עבודה ה WBM-105", "ב18.05 התקבלה חבילת המסירה בסך 900,000 ₪ ממתינים לנספח 5", transfer="טרם",
        ground="צפי עליה לקרקע 11/26"),
    row("X-106", "תחנה ו X-106", "מאחר וההסכם הסתיים והקבלן טרם חתם על ההזמנה התייעצתי", transfer="לא יועבר - פחות מ500 אש\"ח ",
        proc='שו"ש דרום מימוש אופציה'),
    row("X-107", "תחנה ז X-107", "ב13.05 דיון בוועדה. ב07.07 התקבלה הודעת זכייה לקבלן \"בטא\" 287,600 ₪ ממתינים להזמנה",
        proc='התחיל כשו"ש דרום עבר לפומבי', ground=D(2026, 12, 15)),
    row("X-108", "תחנה ח X-108", "אין אישור העירייה לגבי מיקום התחנה.", transfer="טרם"),
    [None] * 10,
    [None, None, "בהקפאה", None, None, None, None, None, None, None],
]
FILL_ORANGE, FILL_GREEN = "FFC000", "92D050"


def fills():
    f = [[None] * 10 for _ in ROWS]
    f[8][1] = FILL_ORANGE        # תחנה ח is coloured "בהקפאה" in the sheet
    f[10][1] = FILL_ORANGE       # the legend row
    return f


def parsed():
    return cs.parse_rows(ROWS, fills(), "סטטוס התקשרויות קבלנים ותיחורים", REPORT)


def by_ident(p, code):
    return next(r for r in p["rows"] if code in r["codes"])


def test_header_date_legend_and_unnamed_name_column():
    p = parsed()
    assert p["as_of"] == "2026-07-09" and p["as_of_stated"]
    assert p["stale_days"] == (REPORT - datetime.date(2026, 7, 9)).days
    assert p["legend"] == {FILL_ORANGE: "בהקפאה"}
    assert p["columns"]["name"] == "שם העבודה"           # the file leaves that header empty
    assert len(p["rows"]) == 8                            # the legend row is not a work


def test_a_plan_is_never_a_step_reached():
    r = by_ident(parsed(), "X-101")
    assert r["stage"] == cs.BIDS
    assert cs.AWARD not in r["proved"]                    # "צפי למכתב זכייה" is a plan
    assert r["next"] == {"what": "מכתב זכייה", "date": "2026-07-15"}
    assert r["runway_days"] == (datetime.date(2026, 9, 30) - REPORT).days


def test_signed_order_contractor_and_award_duration():
    r = by_ident(parsed(), "X-102")
    assert r["group"] == "signed" and r["stage"] == cs.ORDER
    assert r["contractor"] == "קבלן אלפא"
    assert "award_amount" not in r and "1,200,000" not in r["status"] and cs.REDACTED in r["status"]
    assert r["order_no"] == "4001000001"
    assert r["runway_days"] is None and r["next"] is None
    assert r["ground"]["permit_received"] and r["ground"]["permit"] is None
    assert r["days_open"] == (datetime.date(2026, 5, 17) - datetime.date(2026, 2, 5)).days


def test_a_cancelled_tour_is_not_held():
    r = by_ident(parsed(), "X-103")
    assert cs.TOUR not in r["proved"] and r["stage"] == cs.PROCUREMENT


def test_a_postponed_submission_counts_at_its_last_date():
    r = by_ident(parsed(), "X-104")
    assert cs.BIDS not in r["proved"]                     # 06.07 ≤ 9.7, but it moved to 13.7
    assert r["next"] == {"what": "הגשת הצעות", "date": "2026-07-13"}
    assert r["postponed"]


def test_package_estimate_is_not_an_award_and_month_year_ground():
    r = by_ident(parsed(), "WBM-105")
    assert "900,000" not in r["status"] and "package_amount" not in r
    assert r["group"] == "docs" and r["transfer"] is None and "annex5" in r["topics"]
    assert r["ground"]["ground"] == "2026-11-30"          # "11/26" = end of November


def test_agreement_ended_mid_text_is_not_done():
    r = by_ident(parsed(), "X-106")
    assert r["group"] != "signed" and r["outside_procurement"]
    assert "agreement" in r["topics"] and r["family"] == 'שו"ש דרום'


def test_award_without_besach_converted_procedure_and_bare_date():
    r = by_ident(parsed(), "X-107")
    assert r["group"] == "awarded" and r["contractor"] == "בטא"
    assert r["family"] == "פומבי (הוסב)"
    assert r["ground"]["bare"] and r["ground"]["ground"] == "2026-12-15"


def test_frozen_by_the_sheets_own_legend_only():
    p = parsed()
    r = by_ident(p, "X-108")
    assert r["frozen"] and r["group"] == "frozen" and r["marks"] == ["בהקפאה"]
    # an unexplained colour is never read
    f = fills()
    f[1][1] = "FFFF00"
    assert cs.parse_rows(ROWS, f, "", REPORT)["rows"][0]["marks"] == []


def test_events_split_on_dates_and_infer_the_year():
    ev = cs.events("פתיח ב25.01 צעד ראשון .ב03.06 צעד שני", datetime.date(2026, 7, 9))
    assert ev == [{"date": None, "text": "פתיח"},
                  {"date": "2026-01-25", "text": "צעד ראשון"},
                  {"date": "2026-06-03", "text": "צעד שני"}]
    # more than half a year ahead of the writer = last year
    assert cs.events("ב20.12 משהו", datetime.date(2026, 3, 1))[0]["date"] == "2025-12-20"
    # the last update ignores a dated plan still ahead
    r = by_ident(parsed(), "X-101")
    assert r["last_event"] == "2026-07-06"


def test_every_number_counts_rows():
    p = parsed()
    a = cs.analyse(p)
    rows = p["rows"]
    assert a["total"] == len(rows) == sum(a["counts"].values())
    assert a["funnel"][0]["reached"] == len(rows)
    assert sum(s["here"] for s in a["funnel"]) == len(rows)
    for s, nxt in zip(a["funnel"], a["funnel"][1:]):
        assert s["reached"] >= nxt["reached"]
    assert len(a["overdue"]) == sum(1 for r in rows if r["runway_days"] is not None and r["runway_days"] < 0)
    assert sum(f["total"] for f in a["families"]) == len(rows)
    for f in a["families"]:
        assert sum(f["phases"].values()) == f["total"]
        assert sum(f[k] for k, _, _ in cs.GROUPS) == f["total"]
    assert a["stale"]
    overdue_find = next(f for f in a["findings"] if f["filter"] == "overdue")
    assert f"{len(a['overdue'])} עבודות" in overdue_find["text"]


def test_pick_sheet():
    assert cs.pick_sheet(["דוח שבועי עדכני", "סטטוס התקשרויות קבלנים ותיחורים"]) == "סטטוס התקשרויות קבלנים ותיחורים"
    assert cs.pick_sheet(["דוח שבועי עדכני"]) is None


def test_read_file_round_trip(tmp_path):
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    wb.active.title = "דוח שבועי עדכני"
    ws = wb.create_sheet("סטטוס התקשרויות קבלנים ותיחורים")
    for r in ROWS:
        ws.append(r)
    for (ri, ci) in ((9, 2), (11, 2)):
        ws.cell(ri, ci).fill = PatternFill("solid", fgColor="FF" + FILL_ORANGE)
    path = tmp_path / "master.xlsx"
    wb.save(path)
    p = cs.read_file(str(path), REPORT)
    assert len(p["rows"]) == 8 and p["legend"] == {FILL_ORANGE: "בהקפאה"}
    assert by_ident(p, "X-108")["group"] == "frozen"
    assert cs.read_file(str(tmp_path / "x.csv")) is None


def _render(p):
    env = il_format.template_env()
    env.autoescape = True
    a = cs.analyse(p) if p else None
    if p:
        for r in p["rows"]:
            r.update(link=r["codes"][0] if r["codes"] else None, record_fc=None, record_gap_days=None)
    return env.get_template("project_contracts.html").render(
        p=p, a=a, groups=cs.GROUPS, stages=cs.LADDER, topics=cs.TOPICS, soon_days=cs.RUNWAY_SOON_DAYS,
        phase_of=cs.PHASE_OF, group_labels=cs.GROUP_LABELS, stage_labels=cs.STAGE_LABELS, loaded_at_il=None)


def test_page_renders_every_row_and_the_empty_state():
    html = _render(parsed())
    assert html.count('<tbody class="row"') == 8
    assert "/dashboard/projects/p/X-101" in html
    assert "₪" not in html and "1,200,000" not in html and "287,600" not in html
    assert "הגיליון לא עודכן" in html
    assert "עדיין לא נקלט" in _render(None)


def test_page_writes_text_only_and_reads_dates_il():
    src = Path("app/templates/project_contracts.html").read_text(encoding="utf-8")
    assert "innerHTML" not in src
    html = _render(parsed())
    visible = re.sub(r"<[^>]+>", " ", html.split("<script>")[0])
    assert not re.search(r"\b20\d\d-\d\d-\d\d\b", visible)   # people read DD/MM/YYYY


def test_upload_saves_the_sheet_after_a_live_sync_only():
    src = Path("app/routers/projects.py").read_text(encoding="utf-8")
    body = src.split("async def _sync_and_summarize")[1].split("@router")[0]
    live = body.split('if result.get("processed") and not result.get("historical"):')[1]
    assert "contracts_sheet.save_from_file" in live



def test_knowledge_upload_saves_the_sheet_after_a_live_sync_only():
    # Telegram + the files page go through process_master_file, not /projects/upload.
    # It saved the exec summary but never this sheet — the page stayed on 17.10.
    src = Path("app/services/knowledge_service.py").read_text(encoding="utf-8")
    body = src.split("async def process_master_file")[1]
    live = body.split('if sr.get("processed") and not sr.get("historical"):')[1].split("except")[0]
    assert "contracts_sheet.save_from_file" in live

def test_no_money_figure_survives():
    for text in ("ע\"ש קבלן בסך 30,432,546 ש\"ח מס הזמנה", "בסך 872,710", "287,600 ₪",
                 "לא יועבר - פחות מ500 אש\"ח", "בעלות של כ-3 מלש\"ח", "₪ 1,000"):
        out = cs.redact(text)
        assert not re.search(r"\d{3,}|₪", out.replace("4001", "")), out
        assert cs.REDACTED in out
    # dates, order numbers and stage words are not money
    assert cs.redact("ב17.05 התקבל מכתב זכייה הזמנה 4001435775") == "ב17.05 התקבל מכתב זכייה הזמנה 4001435775"


def test_a_payload_stored_before_the_rule_is_cleaned_on_read():
    old = {"rows": [{"status": "התקבל מכתב זכייה בסך 1,200,000 ש\"ח", "award_amount": 1200000,
                     "package_amount": 5, "events": [{"date": "2026-05-17", "text": "בסך 9,000 ₪"}]}],
           "awarded_total": 1200000}
    clean = cs.redact_payload(old)
    assert "awarded_total" not in clean
    row = clean["rows"][0]
    assert "award_amount" not in row and "package_amount" not in row
    assert "1,200,000" not in row["status"] and "9,000" not in row["events"][0]["text"]
    assert row["events"][0]["date"] == "2026-05-17"
