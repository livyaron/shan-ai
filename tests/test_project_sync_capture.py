"""PLAN.md P0 — what the master-file sync captures. Pure functions, no DB."""
from datetime import date

import pytest

from app.services import project_sync as ps

# Real headers of the weekly master file ("דוח שבועי עדכני", Sep 2026 layout).
MASTER_HEADERS = [
    "זיהוי", "פרויקט", "סוג תחנה", 'מנה"פ', 'תו"ב', "יעד תכנית פיתוח", "יעד חשמול מסתמן",
    "סטטוס הפרויקט על ציר הזמן", "זיהוי2", "פירוט שבועי 20/05/2026", "פירוט שבועי 19/8/2026",
    "פירוט שבועי 26/8/26", "פירוט שבועי 02/09/2026", "פירוט סיכונים וחסמים עיקריים", "לטיפול",
    "לטיפול נוסף (במידה ונדרש)", "חוסר במשגיחים", "חוסר בבודקים", "פרויקטים קריטים",
]


def test_column_map_captures_new_fields_without_moving_old_ones():
    m = ps._build_column_map(MASTER_HEADERS)
    by_field = {v: k for k, v in m.items() if v != "__weekly__"}
    assert by_field == {
        "project_identifier": "זיהוי",
        "name": "פרויקט",
        "project_type": "סוג תחנה",
        "manager": 'מנה"פ',
        "controller": 'תו"ב',
        "dev_plan_date": "יעד תכנית פיתוח",
        "estimated_finish_date": "יעד חשמול מסתמן",
        "stage": "סטטוס הפרויקט על ציר הזמן",
        "risks": "פירוט סיכונים וחסמים עיקריים",
        "to_handle": "לטיפול",
        "short_supervisors": "חוסר במשגיחים",
        "short_testers": "חוסר בבודקים",
        "critical_tier": "פרויקטים קריטים",
    }
    assert sum(v == "__weekly__" for v in m.values()) == 4


@pytest.mark.parametrize("name,expected", [
    ("abc_דוח שבועי לסמנכל חטיבת הולכה והשנאה 16.09.2026.xlsx", date(2026, 9, 16)),
    ("abc_דוח שבועי לסמנכל חטיבת הולכה והשנאה 15.4.2026 .xlsx", date(2026, 4, 15)),
    ("abc_דוח שבועי לסמנכל חטיבת הולכה והשנאה 19.8.26.xlsx", date(2026, 8, 19)),
    ("/uploads/0f0f_חוברת1.xlsx", None),           # no date — fall back to the columns
    ("abc_דוח 1.7.xlsx", None),                     # day.month without year is too ambiguous
    ("abc_דוח 31.02.2026.xlsx", None),              # not a real day
])
def test_report_date_from_filename(name, expected):
    assert ps.report_date_from_filename(name) == expected


@pytest.mark.parametrize("col,expected", [
    ("פירוט שבועי 04/03/2026", date(2026, 3, 4)),
    ("פירוט שבועי 04.02.26", date(2026, 2, 4)),
    ("פירוט שבועי 19/8/2026", date(2026, 8, 19)),
    ("פירוט שבועי 26/8/26", date(2026, 8, 26)),
    ("פירוט שבועי 14.1", date(2026, 1, 14)),           # year taken from the report
    ("פירוט שבועי דוח קודם 7.1", date(2026, 1, 7)),
    ("פירוט שבועי", None),
])
def test_weekly_column_date(col, expected):
    assert ps.weekly_column_date(col, date(2026, 3, 25)) == expected


def test_weekly_column_without_year_never_lands_after_the_report():
    # A January report listing "פירוט שבועי 17.12" means last December.
    assert ps.weekly_column_date("פירוט שבועי 17.12", date(2026, 1, 14)) == date(2025, 12, 17)


def test_resolve_report_date_prefers_filename_then_newest_column():
    cols = ["פירוט שבועי 24/06/2026", "פירוט שבועי 01/07/2026"]
    assert ps.resolve_report_date("x_דוח 08.07.2026.xlsx", cols) == date(2026, 7, 8)
    assert ps.resolve_report_date("x_חוברת1.xlsx", cols) == date(2026, 7, 1)


@pytest.mark.parametrize("val,expected", [
    ("01/07/2026 לא אפשרי, יתקבל יעד חדש לאחר קבלת המסמכים",
     (None, "01/07/2026 לא אפשרי, יתקבל יעד חדש לאחר קבלת המסמכים")),
    ("טרם נקבע", (None, "טרם נקבע")),
    ("01/07/2026", (date(2026, 7, 1), None)),        # day-first, not January 7th
    (date(2027, 5, 30), (date(2027, 5, 30), None)),
    (None, (None, None)),
    ("   ", (None, None)),
])
def test_split_finish_date_keeps_prose(val, expected):
    assert ps._split_finish_date(val) == expected


def test_parse_date_rejects_typo_years():
    assert ps._parse_date(date(1, 12, 12)) is None
    assert ps._parse_date(date(2126, 1, 1)) is None
    assert ps._parse_date(date(2026, 12, 31)) == date(2026, 12, 31)


@pytest.mark.parametrize("val,expected", [("כן", True), (" כן ", True), ("לא", False),
                                          (None, None), (float("nan"), None), ("אולי", None)])
def test_parse_bool(val, expected):
    assert ps._parse_bool(val) is expected


def test_text_hash_ignores_whitespace_only_edits():
    assert ps._text_hash("שנאי  הגיע\nלאתר") == ps._text_hash("שנאי הגיע לאתר")
    assert ps._text_hash("שנאי הגיע לאתר") != ps._text_hash("שנאי לא הגיע לאתר")


async def test_draft_sheet_is_never_synced():
    r = await ps.sync_projects_file("/nonexistent.xlsx", sheet_name="דוח שבועי טיוטה")
    assert r["processed"] == 0 and r["errors"] == []


# ── P1: which sheet of a master workbook is the record ────────────────────

def _workbook(tmp_path, sheets: dict[str, list[list]]) -> str:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    path = tmp_path / "דוח 16.09.2026.xlsx"
    wb.save(path)
    return str(path)


def test_pick_master_sheet_skips_the_leading_macro_and_draft_sheets(tmp_path):
    path = _workbook(tmp_path, {
        "מאקרו1": [["x"]],
        "דוח שבועי טיוטה": [["זיהוי", "פרויקט"], ["A-1", "טיוטה"]],
        "גיליון1": [["משגיחים", "בודקים"]],
        "דוח שבועי עדכני": [["זיהוי", "פרויקט"], ["A-1", "אמיתי"]],
    })
    assert ps.pick_master_sheet(path) == "דוח שבועי עדכני"


def test_pick_master_sheet_falls_back_to_the_sheet_with_an_identifier(tmp_path):
    path = _workbook(tmp_path, {
        "רשימות בחירה": [["סטטוס", "סוג"], ["תכנון", "הרחבה"]],
        "Sheet": [["זיהוי", "פרויקט", "מנהל"], ["A-1", "x", "y"]],
    })
    assert ps.pick_master_sheet(path) == "Sheet"


def test_pick_master_sheet_leaves_csv_alone(tmp_path):
    p = tmp_path / "projects.csv"
    p.write_text("זיהוי,פרויקט\nA-1,x\n", encoding="utf-8")
    assert ps.pick_master_sheet(str(p)) is None


def test_file_report_date_orders_a_backfill_even_without_a_dated_name(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "דוח שבועי עדכני"
    ws.append(["זיהוי", "פירוט שבועי 24/06/2026", "פירוט שבועי 01/07/2026"])
    ws.append(["A-1", "a", "b"])
    path = tmp_path / "חוברת1.xlsx"
    wb.save(path)
    assert ps._file_report_date(str(path)) == date(2026, 7, 1)
