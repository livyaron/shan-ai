"""Knowledge-base files: a click opens the file on screen.

PDF goes to the browser's own viewer, DOCX/XLSX are previewed as HTML, a file
missing on disk gets a page (not raw JSON), and the list page links every row
to /open. No DB — a fake session hands back one KnowledgeFile.
"""
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from httpx import ASGITransport, AsyncClient

from app.database import get_db_session
from app.main import app
from app.models import KnowledgeFile, User
from app.routers import files as files_router
from app.routers.login import get_current_user
from app.services import file_preview


def _docx(path: Path) -> None:
    from docx import Document
    d = Document()
    d.add_heading("נוהל פיקוח", level=1)
    d.add_paragraph("שורה ראשונה של הנוהל")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "שלב", "אחראי"
    t.cell(1, 0).text, t.cell(1, 1).text = "בדיקות", "מנה\"פ"
    d.save(path)


def _xlsx(path: Path, rows: int = 3) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "דוח שבועי עדכני"
    ws.append(["זיהוי", "שם", "יעד"])
    for i in range(rows):
        ws.append([f"P{i}", f"תחנה {i}", datetime(2026, 9, 16)])
    wb.create_sheet("ריק")
    wb.save(path)


def test_preview_docx_keeps_order_headings_and_tables(tmp_path):
    p = tmp_path / "a.docx"
    _docx(p)
    blocks = file_preview.preview_docx(str(p))
    assert [b["kind"] for b in blocks] == ["h", "p", "table"]
    assert blocks[0]["text"] == "נוהל פיקוח"
    assert blocks[2]["rows"] == [["שלב", "אחראי"], ["בדיקות", "מנה\"פ"]]


def test_preview_xlsx_reads_sheets_dates_and_caps_rows(tmp_path, monkeypatch):
    p = tmp_path / "a.xlsx"
    _xlsx(p, rows=10)
    monkeypatch.setattr(file_preview, "MAX_ROWS", 5)
    sheets = file_preview.preview_xlsx(str(p))
    assert [s["name"] for s in sheets] == ["דוח שבועי עדכני", "ריק"]
    first = sheets[0]
    assert first["rows"][0] == ["זיהוי", "שם", "יעד"]
    assert first["rows"][1] == ["P0", "תחנה 0", "16/09/2026"]
    assert len(first["rows"]) == 5 and first["truncated_rows"] is True
    assert sheets[1]["rows"] == []


class _Session:
    def __init__(self, kf):
        self.kf = kf

    async def get(self, _model, _id):
        return self.kf


async def _get(kf, url, monkeypatch):
    async def fake_user():
        return User(id=1, username="u", password_hash="", is_admin=True)

    async def fake_session():
        yield _Session(kf)

    async def no_pending(*_a, **_k):
        return 0

    monkeypatch.setattr(files_router, "_pending_approvals_count", no_pending)
    app.dependency_overrides[get_current_user] = fake_user
    app.dependency_overrides[get_db_session] = fake_session
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            return await c.get(url)
    finally:
        app.dependency_overrides.clear()


def _kf(path: Path, ftype: str, name: str) -> KnowledgeFile:
    return KnowledgeFile(id=7, original_name=name, file_path=str(path), file_type=ftype,
                         file_size=1234, status="ready", is_master=False,
                         created_at=datetime(2026, 9, 16, 8, 0))


@pytest.mark.asyncio
async def test_open_docx_and_xlsx_render_on_screen(tmp_path, monkeypatch):
    d = tmp_path / "x.docx"
    _docx(d)
    r = await _get(_kf(d, "docx", "נוהל.docx"), "/dashboard/files/7/open", monkeypatch)
    assert r.status_code == 200
    assert "שורה ראשונה של הנוהל" in r.text and '<article class="doc"' in r.text

    x = tmp_path / "x.xlsx"
    _xlsx(x)
    r = await _get(_kf(x, "xlsx", "מאסטר.xlsx"), "/dashboard/files/7/open", monkeypatch)
    assert r.status_code == 200
    assert "תחנה 2" in r.text and "דוח שבועי עדכני" in r.text


@pytest.mark.asyncio
async def test_open_pdf_embeds_inline_view(tmp_path, monkeypatch):
    p = tmp_path / "x.pdf"
    p.write_bytes(b"%PDF-1.4\n%%EOF\n")
    kf = _kf(p, "pdf", "נוהל.pdf")
    r = await _get(kf, "/dashboard/files/7/open", monkeypatch)
    assert r.status_code == 200
    assert '<iframe class="pdf-frame" src="/dashboard/files/7/view"' in r.text

    r = await _get(kf, "/dashboard/files/7/view", monkeypatch)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"].startswith("inline")

    r = await _get(kf, "/dashboard/files/7/view?download=1", monkeypatch)
    assert r.headers["content-disposition"].startswith("attachment")


@pytest.mark.asyncio
async def test_open_missing_file_is_a_page_not_json(tmp_path, monkeypatch):
    r = await _get(_kf(tmp_path / "gone.pdf", "pdf", "gone.pdf"), "/dashboard/files/7/open", monkeypatch)
    assert r.status_code == 200
    assert "הקובץ המקורי לא נמצא בשרת" in r.text
    assert "<iframe" not in r.text


def test_list_page_links_every_file_to_open():
    tpl = files_router.templates.env.get_template("files.html")
    html = tpl.render(
        request=SimpleNamespace(url=SimpleNamespace(path="/dashboard/files")),
        current_user=SimpleNamespace(id=1, username="u", is_admin=True, role=None),
        files=[dict(id=3, original_name="it's.pdf", file_type="pdf", file_size=10, uploader_name="u",
                    summary="", chunk_count=0, status="ready", is_master=False, created_at="01/01/2026 00:00")],
        msg=None, error=None, pending_approvals=0,
    )
    assert 'href="/dashboard/files/3/open"' in html
    assert 'data-open="/dashboard/files/3/open"' in html
    # the old Bootstrap table is what painted rows black on hover
    assert "table-hover" not in html
    # a quote in the name must not break the delete confirm's JS string
    assert "confirm('למחוק את it's.pdf" not in html


def test_column_letters_match_excel():
    assert file_preview.COL_LETTERS[:3] == ["A", "B", "C"]
    assert file_preview._col_letter(25) == "Z" and file_preview._col_letter(26) == "AA"
    assert len(file_preview.COL_LETTERS) == file_preview.MAX_COLS
