"""The "תקציר הנהלה" sheet of the master workbook → the projects dashboard.

The PM department writes a management summary inside the weekly master file
(KPI pairs, a priorities table, proposed messages, a shortage table). The
dashboard shows THAT, as written — it is the division's own reading of the
week, and the number of KPIs and rows changes from file to file, so nothing
here assumes a fixed layout:

    a row with one text cell            → title / subtitle / section title
    a row of (label, number) pairs      → KPIs
    the first all-short-text row of a section → its header
    anything else                       → a row of the current section

`parse_rows` is pure (tested on synthetic rows); `read_file` is the only file
access; `save` / `latest` the only DB access, both best-effort.
"""
import datetime
import logging
import re

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

SHEET_MARKER = "תקציר הנהלה"
HEADER_CELL_CHARS = 40        # a header cell is a short label, never a sentence
CELL_CHARS = 600              # one cell on screen; the file keeps the rest
ID_HEADERS = ("מזהה", "זיהוי")  # this column links to the project page

_SHEET_DATE = re.compile(r"(\d{1,2})[./](\d{1,2})[./](\d{2,4})")
_NUMERIC = re.compile(r"^-?\d+(\.\d+)?$")


def sheet_date(name: str) -> datetime.date | None:
    m = _SHEET_DATE.search(name or "")
    if not m:
        return None
    d, mo, y = (int(g) for g in m.groups())
    y = y + 2000 if y < 100 else y
    try:
        return datetime.date(y, mo, d)
    except ValueError:
        return None


def pick_sheet(names: list[str]) -> str | None:
    """The summary sheet; with several dated copies, the newest date wins."""
    found = [n for n in names if SHEET_MARKER in n]
    if not found:
        return None
    return max(found, key=lambda n: (sheet_date(n) or datetime.date.min, found.index(n)))


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _numeric_text(v) -> bool:
    return _is_num(v) or (isinstance(v, str) and bool(_NUMERIC.match(v.strip())))


def _cell(v) -> str:
    """A table cell as text. A bare 0 is an empty source cell seen through a formula."""
    if v is None or (_is_num(v) and v == 0) or (isinstance(v, str) and v.strip() == "0"):
        return ""
    if isinstance(v, datetime.datetime):
        v = v.date()
    if isinstance(v, datetime.date):
        return v.strftime("%d/%m/%Y")
    if _is_num(v):
        return str(int(v)) if float(v).is_integer() else str(v)
    s = str(v).strip()
    return s if len(s) <= CELL_CHARS else s[:CELL_CHARS - 1] + "…"


def _kpi_pairs(cells: list) -> list[list] | None:
    """[label, number, label, number…] → pairs; None when the row is not that."""
    vals = [v for v in cells if v is not None and str(v).strip() != ""]
    if not vals or len(vals) % 2:
        return None
    pairs = []
    for label, value in zip(vals[::2], vals[1::2]):
        if not isinstance(label, str) or _numeric_text(label) or not _is_num(value):
            return None
        pairs.append([label.strip(), int(value) if float(value).is_integer() else value])
    return pairs


def _is_header(cells: list) -> bool:
    vals = [v for v in cells if v is not None and str(v).strip() != ""]
    return len(vals) >= 2 and all(
        isinstance(v, str) and not _numeric_text(v) and len(v.strip()) <= HEADER_CELL_CHARS
        for v in vals)


def parse_rows(rows: list[list], sheet: str = "") -> dict:
    """Rows of raw cell values (column order kept) → the dashboard payload."""
    out = {"sheet": sheet, "as_of": None, "title": None, "subtitle": None,
           "kpis": [], "sections": []}
    d = sheet_date(sheet)
    out["as_of"] = d.isoformat() if d else None
    section = None
    for raw in rows:
        cells = list(raw)
        while cells and (cells[-1] is None or str(cells[-1]).strip() == ""):
            cells.pop()
        filled = [v for v in cells if v is not None and str(v).strip() != ""]
        if not filled:
            continue
        if len(filled) == 1 and isinstance(filled[0], str) and not _numeric_text(filled[0]):
            text = filled[0].strip()
            if out["title"] is None and not out["kpis"] and not out["sections"]:
                out["title"] = text
            elif out["subtitle"] is None and not out["kpis"] and not out["sections"]:
                out["subtitle"] = text
            else:
                section = {"title": text, "header": None, "rows": [], "id_col": None}
                out["sections"].append(section)
            continue
        pairs = _kpi_pairs(cells)
        if pairs and section is None:
            out["kpis"].extend(pairs)
            continue
        if section is None:
            section = {"title": None, "header": None, "rows": [], "id_col": None}
            out["sections"].append(section)
        if section["header"] is None and not section["rows"] and _is_header(cells):
            section["header"] = [_cell(v) for v in cells]
            section["id_col"] = next(
                (i for i, h in enumerate(section["header"]) if h.strip() in ID_HEADERS), None)
            continue
        section["rows"].append([_cell(v) for v in cells])

    # every row of a section is as wide as its widest row, so the table lines up
    for s in out["sections"]:
        width = max([len(s["header"] or [])] + [len(r) for r in s["rows"]])
        if s["header"] is not None:
            s["header"] += [""] * (width - len(s["header"]))
        for r in s["rows"]:
            r += [""] * (width - len(r))
        # a column that is empty in every row (a spacer) is not drawn
        keep = [i for i in range(width)
                if any(r[i] for r in s["rows"]) or (s["header"] and s["header"][i])]
        if s["header"] is not None:
            s["header"] = [s["header"][i] for i in keep]
        s["rows"] = [[r[i] for i in keep] for r in s["rows"]]
        if s["id_col"] is not None:
            s["id_col"] = keep.index(s["id_col"]) if s["id_col"] in keep else None
    out["sections"] = [s for s in out["sections"] if s["rows"]]
    return out


def read_file(file_path: str) -> dict | None:
    """The summary sheet of an XLSX, parsed; None when the workbook has none."""
    if not str(file_path).lower().endswith((".xlsx", ".xlsm")):
        return None
    import openpyxl
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:
        logger.warning(f"exec_summary: cannot open {file_path}: {exc}")
        return None
    try:
        name = pick_sheet(wb.sheetnames)
        if name is None:
            logger.info(f"exec_summary: no '{SHEET_MARKER}' sheet in {file_path} — dashboard keeps the previous one")
            return None
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    finally:
        wb.close()
    payload = parse_rows(rows, name)
    if not payload["kpis"] and not payload["sections"]:
        logger.warning(f"exec_summary: sheet '{name}' parsed empty")
        return None
    return payload


async def save(session: AsyncSession, payload: dict, source_file: str | None) -> None:
    from app.models import ProjectExecSummary
    try:
        session.add(ProjectExecSummary(
            sheet_name=payload.get("sheet") or "",
            as_of=datetime.date.fromisoformat(payload["as_of"]) if payload.get("as_of") else None,
            source_file=(source_file or "")[:300],
            payload=payload,
        ))
        await session.commit()
        logger.info(f"exec_summary: saved '{payload.get('sheet')}' — {len(payload.get('kpis') or [])} KPIs, "
                    f"{len(payload.get('sections') or [])} tables")
    except Exception:
        logger.exception("exec_summary: save failed")
        await session.rollback()


async def save_from_file(file_path: str, source_file: str | None) -> None:
    """Parse + save in a session of its own — called after a LIVE sync only. Never raises."""
    import asyncio
    from app.database import async_session_maker
    try:
        payload = await asyncio.get_running_loop().run_in_executor(None, read_file, file_path)
        if payload:
            async with async_session_maker() as session:
                await save(session, payload, source_file)
    except Exception:
        logger.exception("exec_summary: save_from_file failed")


async def latest(session: AsyncSession) -> dict | None:
    """The summary of the newest LIVE master upload (history replays never save one)."""
    from app.models import ProjectExecSummary
    row = (await session.execute(
        select(ProjectExecSummary.payload, ProjectExecSummary.created_at)
        .order_by(ProjectExecSummary.id.desc()).limit(1)
    )).first()
    if not row:
        return None
    payload, created = row
    from app.services.il_format import il_date, il_datetime
    return {**payload,
            "as_of_il": il_date(payload.get("as_of")) if payload.get("as_of") else None,
            "loaded_at_il": il_datetime(created) if created else None}
