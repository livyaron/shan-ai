"""קישור פרויקטים — the completion table as an Excel file (user decision 2026-10-19, option ב).

Export: every item the work list says only a person can give, one row per
item, with a drop-down of the legal values. Import: the same file back →
validate → preview → apply. Pure functions here (build / parse / validate,
tested without a DB); the async apply lives in lessons_project_link.

Safety: a row is matched by its hidden key (file identifier / lessons project
id / the file's manager name — never by a project name); an empty cell changes
nothing; a row whose "current value" no longer matches what the system holds
now is refused as stale; an item that is no longer open is skipped.
"""
from __future__ import annotations

import io
import re
from typing import Any

SHEET_ITEMS = "להשלמה"
SHEET_LISTS = "רשימות"
HEADERS = ("מפתח", "מזהה פרויקט", "שם הפרויקט", "מנהל פרויקט", "שדה להשלמה", "ערך קיים", "ערך להשלמה",
           "הערות / סיבה")
COL = {h: i for i, h in enumerate(HEADERS)}
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024

FIELD_LABELS = {
    "stage": "שלב",
    "project_type": "סוג פרויקט",
    "alias": "משתמש Shan-AI לשם הזה (כל הפרויקטים שלו)",
    "pm": "מנה״פ לפרויקט",
}
STATUS_LABELS = {
    "ok": "יעודכן",
    "closed": "דולג — כבר לא פתוח",
    "stale": "נדחה — הערך הקיים השתנה מאז הייצוא",
    "invalid": "נדחה — ערך לא חוקי",
    "duplicate": "נדחה — השורה מופיעה פעמיים",
}

_USER_ID = re.compile(r"#(\d+)\s*$")


def user_option(u: dict[str, Any]) -> str:
    """How a user appears in the drop-down: the id is what is read back, never the name."""
    return f"{(u.get('username') or '').strip()} · #{u['id']}"


def build_workbook(items: list[dict[str, Any]], stages: list[str], types: dict[str, str],
                   users: list[dict[str, Any]]) -> bytes:
    """The completion file. `items`: dicts with key, kind, identifier, name, manager, current, note."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_ITEMS
    ws.sheet_view.rightToLeft = True
    ws.append(list(HEADERS))
    for c in ws[1]:
        c.font = Font(bold=True)
    lists = wb.create_sheet(SHEET_LISTS)
    lists.sheet_state = "hidden"
    columns = {"stage": list(stages), "project_type": list(types.values()),
               "user": [user_option(u) for u in users]}
    ranges = {}
    for col_i, (name, values) in enumerate(columns.items(), start=1):
        letter = chr(ord("A") + col_i - 1)
        lists.cell(row=1, column=col_i, value=name)
        for r, v in enumerate(values, start=2):
            lists.cell(row=r, column=col_i, value=v)
        ranges[name] = f"'{SHEET_LISTS}'!${letter}$2:${letter}${max(2, len(values) + 1)}"
    validations = {k: DataValidation(type="list", formula1=f"={v}", allow_blank=True, showErrorMessage=True,
                                     errorTitle="ערך לא חוקי", error="בחר ערך מהרשימה")
                   for k, v in ranges.items()}
    for dv in validations.values():
        ws.add_data_validation(dv)
    fill = PatternFill("solid", fgColor="FFF4CC")
    for i, it in enumerate(items, start=2):
        ws.append([it["key"], it.get("identifier") or "", it.get("name") or "", it.get("manager") or "",
                   FIELD_LABELS[it["kind"]], it.get("current") or "", "", it.get("note") or ""])
        cell = ws.cell(row=i, column=COL["ערך להשלמה"] + 1)
        cell.fill = fill
        validations["user" if it["kind"] in ("alias", "pm") else it["kind"]].add(cell.coordinate)
    ws.column_dimensions["A"].hidden = True
    for letter, width in zip("BCDEFGH", (16, 40, 22, 30, 22, 30, 40), strict=True):
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.alignment = Alignment(horizontal="right")
    ws.freeze_panes = "B2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class SheetError(Exception):
    """The file is not the completion table (shown to the admin as is)."""


def parse_workbook(data: bytes) -> list[dict[str, Any]]:
    """Rows of the completion table: key, current, value (as typed), row number."""
    from openpyxl import load_workbook

    if len(data) > MAX_BYTES:
        raise SheetError("הקובץ גדול מדי")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise SheetError("לא ניתן לקרוא את הקובץ — האם זה קובץ Excel שהורד מכאן?") from e
    if SHEET_ITEMS not in wb.sheetnames:
        raise SheetError(f"אין בקובץ גיליון בשם \"{SHEET_ITEMS}\" — יש להעלות את הקובץ שהורד מכאן")
    rows = list(wb[SHEET_ITEMS].iter_rows(values_only=True))
    if not rows or tuple(str(h or "").strip() for h in rows[0][:len(HEADERS)]) != HEADERS:
        raise SheetError("כותרות העמודות שונו — יש להעלות את הקובץ כפי שהורד, רק עם עמודת \"ערך להשלמה\" ממולאת")
    if len(rows) - 1 > MAX_ROWS:
        raise SheetError("יותר מדי שורות")
    out = []
    for n, r in enumerate(rows[1:], start=2):
        r = tuple(r) + (None,) * (len(HEADERS) - len(r))
        key = str(r[COL["מפתח"]] or "").strip()
        if not key:
            continue
        out.append({"row": n, "key": key,
                    "current": str(r[COL["ערך קיים"]] or "").strip(),
                    "value": str(r[COL["ערך להשלמה"]] or "").strip(),
                    "label": " · ".join(str(x).strip() for x in (r[COL["מזהה פרויקט"]], r[COL["שם הפרויקט"]],
                                                                 r[COL["שדה להשלמה"]]) if x)})
    return out


def parse_value(kind: str, raw: str, stages: list[str], types: dict[str, str],
                user_ids: set[int]) -> Any | None:
    """The value as the system stores it, or None when it is not a legal one. Pure."""
    if kind == "stage":
        return raw if raw in stages else None
    if kind == "project_type":
        by_label = {v: k for k, v in types.items()}
        return by_label.get(raw) or (raw if raw in types else None)
    m = _USER_ID.search(raw)
    if m and int(m.group(1)) in user_ids:
        return int(m.group(1))
    return None


def validate(rows: list[dict[str, Any]], open_items: dict[str, dict[str, Any]], stages: list[str],
             types: dict[str, str], user_ids: set[int]) -> dict[str, Any]:
    """The preview: each filled row's status, and the changes to apply. Pure.
    Empty value = nothing to do (counted, not listed)."""
    seen: set[str] = set()
    out, apply = [], []
    empty = 0
    for r in rows:
        if not r["value"]:
            empty += 1
            continue
        item = open_items.get(r["key"])
        if r["key"] in seen:
            status = "duplicate"
        elif item is None:
            status = "closed"
        elif (item.get("current") or "").strip() != r["current"]:
            status = "stale"
        else:
            parsed = parse_value(item["kind"], r["value"], stages, types, user_ids)
            status = "ok" if parsed is not None else "invalid"
            if parsed is not None:
                apply.append({**item, "value": parsed, "row": r["row"]})
        seen.add(r["key"])
        out.append({"row": r["row"], "label": r["label"], "value": r["value"], "status": status,
                    "status_label": STATUS_LABELS[status]})
    counts = {s: sum(1 for o in out if o["status"] == s) for s in STATUS_LABELS}
    return {"rows": out, "apply": apply, "counts": {**counts, "empty": empty}}
